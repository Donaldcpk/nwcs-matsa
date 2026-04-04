#!/usr/bin/env python3
"""
Sync RPG Maker MZ data/Enemies.json from enermy STAT.xlsx (or similar).

Requires: pip install openpyxl

Default xlsx: ~/Downloads/enermy STAT.xlsx
Default enemies: <repo>/data/Enemies.json

Usage:
  python3 tools/sync_enemies_from_xlsx.py              # dry-run
  python3 tools/sync_enemies_from_xlsx.py --write        # backup + write

MZ enemy params index: 0 MHP, 1 MMP, 2 ATK, 3 DEF, 4 MAT, 5 MDF, 6 AGI, 7 LUK
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    print("Missing openpyxl. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

# 試算表欄名（strip 後）→ params 索引 或 特殊鍵
HEADER_TO_TARGET: dict[str, tuple[str, int] | str] = {
    "最大hp": ("param", 0),
    "最大mp": ("param", 1),
    "攻擊": ("param", 2),
    "防禦": ("param", 3),
    "魔法攻擊": ("param", 4),
    "魔法防禦": ("param", 5),
    "敏捷度": ("param", 6),
    "運氣": ("param", 7),
    "經驗值": "exp",
    "金幣": "gold",
}


def repo_root() -> Path:
    return Path(__file__).resolve().parent.parent


def norm_header(h: Any) -> str:
    if h is None:
        return ""
    return str(h).strip()


def parse_enemy_id(cell: Any) -> int | None:
    if cell is None:
        return None
    if isinstance(cell, bool):
        return None
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float):
        return int(round(cell))
    s = str(cell).strip()
    if not s:
        return None
    low = s.lower()
    if low.startswith("id"):
        s = s[2:].lstrip()
    m = re.match(r"^(\d+)$", s)
    if m:
        return int(m.group(1))
    return None


def cell_to_int(cell: Any) -> int | None:
    """None/空白不覆寫；數字轉 int。"""
    if cell is None:
        return None
    if isinstance(cell, str) and not cell.strip():
        return None
    if isinstance(cell, bool):
        return int(cell)
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float):
        return int(round(cell))
    try:
        return int(round(float(str(cell).strip())))
    except (TypeError, ValueError):
        return None


def build_column_targets(headers: tuple[Any, ...]) -> list[str | tuple[str, int] | None]:
    """每欄對應 'id' / ('param', i) / 'exp' / 'gold' / None（略過）。"""
    out: list[str | tuple[str, int] | None] = []
    for i, h in enumerate(headers):
        hn = norm_header(h)
        if i == 0:
            out.append("id")
            continue
        if hn.lower().startswith("id") and hn[2:].strip().isdigit():
            # 標題列誤標成 id10 這類仍當第一欄處理過了
            out.append("id")
            continue
        if hn in HEADER_TO_TARGET:
            out.append(HEADER_TO_TARGET[hn])
            continue
        out.append(None)
    return out


def load_rows(
    xlsx: Path, sheet: str | None
) -> tuple[tuple[Any, ...], list[tuple[Any, ...]]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if sheet:
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise ValueError("empty sheet")
        data_rows = [r for r in rows_iter if r and r[0] is not None]
        return header, data_rows
    finally:
        wb.close()


def build_id_index(enemies: list[Any]) -> dict[int, int]:
    idx: dict[int, int] = {}
    for i, e in enumerate(enemies):
        if isinstance(e, dict) and isinstance(e.get("id"), int):
            idx[e["id"]] = i
    return idx


def write_enemies_json(path: Path, data: list[Any]) -> None:
    """與現有 RMMZ Enemies.json 相同排版：[\n null,\n {...},\n ...\n]"""
    with path.open("w", encoding="utf-8", newline="\n") as f:
        f.write("[\n")
        f.write("null,\n")
        last = len(data) - 1
        for i in range(1, len(data)):
            line = json.dumps(data[i], ensure_ascii=False, separators=(",", ":"))
            if i < last:
                f.write(line + ",\n")
            else:
                f.write(line + "\n")
        f.write("]\n")


def main() -> int:
    default_xlsx = Path.home() / "Downloads" / "enermy STAT.xlsx"
    ap = argparse.ArgumentParser(description="Sync Enemies.json from Excel stats.")
    ap.add_argument("--xlsx", type=Path, default=default_xlsx, help="Path to .xlsx")
    ap.add_argument(
        "--enemies",
        type=Path,
        default=repo_root() / "data" / "Enemies.json",
        help="Path to Enemies.json",
    )
    ap.add_argument("--sheet", default=None, help="Worksheet name (default: first sheet)")
    ap.add_argument(
        "--write",
        action="store_true",
        help="Write Enemies.json (default: dry-run only)",
    )
    ap.add_argument(
        "--no-backup",
        action="store_true",
        help="With --write, skip copying Enemies.json.bak",
    )
    args = ap.parse_args()

    if not args.xlsx.is_file():
        print(f"Missing xlsx: {args.xlsx}", file=sys.stderr)
        return 1
    if not args.enemies.is_file():
        print(f"Missing Enemies.json: {args.enemies}", file=sys.stderr)
        return 1

    header, data_rows = load_rows(args.xlsx, args.sheet)
    col_targets = build_column_targets(header)

    print("Sheet:", args.sheet or "(first)")
    print("Headers:", header)
    unmapped = [
        (i, norm_header(header[i]))
        for i, t in enumerate(col_targets)
        if t is None and norm_header(header[i])
    ]
    if unmapped:
        print("Warning: unmapped columns (skipped):", unmapped)

    with args.enemies.open(encoding="utf-8") as f:
        enemies: list[Any] = json.load(f)

    working = copy.deepcopy(enemies)
    id_index = build_id_index(working)

    missing_ids: list[int] = []
    updated = 0
    skipped_no_id = 0

    for row in data_rows:
        eid = None
        param_deltas: dict[int, int] = {}
        exp_val: int | None = None
        gold_val: int | None = None

        for col_idx, target in enumerate(col_targets):
            if col_idx >= len(row):
                break
            cell = row[col_idx]
            if target == "id":
                eid = parse_enemy_id(cell)
                continue
            if target is None:
                continue
            val = cell_to_int(cell)
            if val is None:
                continue
            if target == "exp":
                exp_val = val
            elif target == "gold":
                gold_val = val
            elif isinstance(target, tuple) and target[0] == "param":
                param_deltas[target[1]] = val

        if eid is None:
            skipped_no_id += 1
            continue
        idx = id_index.get(eid)
        if idx is None:
            missing_ids.append(eid)
            continue

        ent = working[idx]
        assert isinstance(ent, dict)

        changed = False
        if param_deltas:
            params = list(ent.get("params") or [0] * 8)
            while len(params) < 8:
                params.append(0)
            for pi, v in param_deltas.items():
                if 0 <= pi < 8 and params[pi] != v:
                    params[pi] = v
                    changed = True
            if changed:
                ent["params"] = params
        if exp_val is not None and ent.get("exp") != exp_val:
            ent["exp"] = exp_val
            changed = True
        if gold_val is not None and ent.get("gold") != gold_val:
            ent["gold"] = gold_val
            changed = True

        if changed:
            updated += 1

    print(
        f"Rows in xlsx (non-empty id col): {len(data_rows)}, "
        f"enemies updated: {updated}, missing id in JSON: {len(missing_ids)}, "
        f"skipped (no id): {skipped_no_id}"
    )
    if missing_ids:
        print("Missing enemy ids (first 30):", missing_ids[:30])

    if not args.write:
        print("Dry-run only. Pass --write to save (creates Enemies.json.bak unless --no-backup).")
        return 0

    bak = args.enemies.with_suffix(args.enemies.suffix + ".bak")
    if not args.no_backup:
        shutil.copy2(args.enemies, bak)
        print(f"Backup: {bak}")

    write_enemies_json(args.enemies, working)
    print(f"Wrote: {args.enemies}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
